import os
import shutil
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

DOCX_PATH = r"C:\Users\parth\Downloads\TravelBrain_Capstone_PhaseI_Report.docx"
BACKUP_PATH = r"C:\Users\parth\Downloads\TravelBrain_Capstone_PhaseI_Report_Backup.docx"
DIAGRAMS_DIR = r"C:\Users\parth\.gemini\antigravity-cli\brain\4bee6099-ee90-43d3-a590-76a48053dbb2\diagrams"

def update_docx():
    if not os.path.exists(BACKUP_PATH):
        shutil.copyfile(DOCX_PATH, BACKUP_PATH)
        print("Backup created at:", BACKUP_PATH)

    doc = Document(DOCX_PATH)

    # Styles
    navy = RGBColor(20, 50, 125)

    # Find paragraphs with figure placeholders and insert images
    fig1_path = os.path.join(DIAGRAMS_DIR, "fig1_workflow_comparison.png")
    fig2_path = os.path.join(DIAGRAMS_DIR, "fig2_rag_vs_multiagent.png")
    fig3_path = os.path.join(DIAGRAMS_DIR, "fig3_master_architecture.png")
    fig4_path = os.path.join(DIAGRAMS_DIR, "fig4_langgraph_flow.png")
    fig5_path = os.path.join(DIAGRAMS_DIR, "fig5_mcp_topology.png")

    inserted = set()

    for p in list(doc.paragraphs):
        text = p.text.strip()
        
        # Check Fig 6.0 / Manual vs TravelBrain in Chapter 1
        if "Fig 6.0: Manual planning vs. TravelBrain" in text and "fig1" not in inserted:
            print("Found Fig 6.0 placeholder, inserting workflow comparison...")
            p_img = p.insert_paragraph_before()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.add_run().add_picture(fig1_path, width=Inches(6.2))
            p.text = "Figure 1.0: Manual Planning Bottlenecks vs. TravelBrain Autonomous Pipeline"
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            inserted.add("fig1")

        # Check Fig 1.0: Overall System Architecture in Chapter 4
        elif "Fig 1.0: Overall System Architecture" in text and "fig3" not in inserted:
            print("Found Fig 1.0 placeholder, inserting 5-layer master architecture...")
            p_img = p.insert_paragraph_before()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.add_run().add_picture(fig3_path, width=Inches(6.2))
            p.text = "Figure 2.0: Comprehensive 5-Layer Master Architecture of TravelBrain"
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            inserted.add("fig3")

        # Check Fig 4.0: Multi-Agent Sequential Workflow in Chapter 5
        elif "Fig 4.0: Multi-Agent Sequential Workflow" in text and "fig4" not in inserted:
            print("Found Fig 4.0 placeholder, inserting LangGraph flow...")
            p_img = p.insert_paragraph_before()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.add_run().add_picture(fig4_path, width=Inches(6.2))
            p.text = "Figure 3.0: LangGraph Sequential Control Flow & Shared State Transitions"
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            inserted.add("fig4")

        # Check Fig 5.0: MCP Integration Diagram in Chapter 5
        elif "Fig 5.0: MCP (Model Context Protocol)" in text and "fig5" not in inserted:
            print("Found Fig 5.0 placeholder, inserting MCP topology...")
            p_img = p.insert_paragraph_before()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.add_run().add_picture(fig5_path, width=Inches(6.2))
            p.text = "Figure 4.0: Model Context Protocol (MCP) Standardized Tool Integration Topology"
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            inserted.add("fig5")

        # Check Fig A.4: RAG vs multi-agent in Appendix
        elif "Fig A.4: RAG vs. multi-agent" in text and "fig2" not in inserted:
            print("Found Fig A.4 placeholder, inserting RAG vs Multi-Agent diagram...")
            p_img = p.insert_paragraph_before()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.add_run().add_picture(fig2_path, width=Inches(6.2))
            p.text = "Figure A.4: Conceptual & Architectural Comparison: Traditional Vector RAG vs. TravelBrain Multi-Agent Engine"
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            inserted.add("fig2")

    doc.save(DOCX_PATH)
    print("Successfully updated DOCX with high-res diagrams at:", DOCX_PATH)

if __name__ == "__main__":
    update_docx()
